import os
import re
import time
import openpyxl
from playwright.sync_api import sync_playwright

EXCEL_PATH = r"D:\Anomali\Anomaliku.xlsx"
USER_DATA_DIR = r"D:\Anomali\chrome-sso-profile"

def clean_name(name):
    if not name:
        return ""
    cleaned = re.sub(r"[<>()[\]]", " ", str(name))
    cleaned = re.sub(r"\b(UTP|UPT|PETANI)\b", " ", cleaned, flags=re.IGNORECASE)
    parts = [p.strip() for p in cleaned.split() if len(p.strip()) > 2]
    return parts[-1] if parts else str(name).strip()

def close_stale_chrome():
    """Tutup proses Chrome sisa yang masih mengunci profil chrome-sso-profile"""
    try:
        import subprocess
        cmd = 'Get-CimInstance Win32_Process -Filter "name = \'chrome.exe\'" | Where-Object { $_.CommandLine -like "*chrome-sso-profile*" } | ForEach-Object { Stop-Process -Id $_.ProcessId -Force }'
        subprocess.run(["powershell", "-Command", cmd], capture_output=True, timeout=10)
        time.sleep(1.5)
    except Exception:
        pass

def run_single_test():
    if not os.path.exists(EXCEL_PATH):
        print(f"[ERROR] File {EXCEL_PATH} tidak ditemukan.")
        return

    wb = openpyxl.load_workbook(EXCEL_PATH)
    ws = wb.active

    headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
    if "Status_Submit_FASIH" not in headers:
        col_submit = ws.max_column + 1
        ws.cell(row=1, column=col_submit, value="Status_Submit_FASIH")
    else:
        col_submit = headers.index("Status_Submit_FASIH") + 1

    col_status = headers.index("Status_Hasil_Cek") + 1 if "Status_Hasil_Cek" in headers else 23
    col_url = headers.index("Link Edit FASIH") + 1 if "Link Edit FASIH" in headers else 22
    col_nama = headers.index("Nama Usaha") + 1 if "Nama Usaha" in headers else 13

    # Cari 1 sampel pertama yang Status_Hasil_Cek == '#DC_00' dan Status_Submit_FASIH belum 'TERKIRIM'
    target_row = None
    for r in range(2, ws.max_row + 1):
        status_val = ws.cell(row=r, column=col_status).value
        submit_val = ws.cell(row=r, column=col_submit).value
        url_val = ws.cell(row=r, column=col_url).value
        nama_val = ws.cell(row=r, column=col_nama).value

        if status_val == "#DC_00" and submit_val != "TERKIRIM" and url_val and "http" in str(url_val):
            target_row = {
                "row_idx": r,
                "nama_usaha": str(nama_val or "").strip(),
                "url": str(url_val).strip(),
                "status": str(status_val).strip()
            }
            break

    if not target_row:
        print("[INFO] Tidak ditemukan baris sampel #DC_00 yang belum berstatus 'TERKIRIM'.")
        return

    print("==================================================")
    print("   UJI COBA 1 SAMPEL REMARK FASIH-SM")
    print(f"   Baris Excel : {target_row['row_idx']}")
    print(f"   Nama Usaha  : {target_row['nama_usaha']}")
    print(f"   Remark Kode : {target_row['status']}")
    print(f"   URL Target  : {target_row['url']}")
    print("==================================================")

    close_stale_chrome()
    with sync_playwright() as p:
        print("\n[1/6] Membuka Google Chrome di layar dengan profil SSO...")
        context = p.chromium.launch_persistent_context(
            user_data_dir=USER_DATA_DIR,
            channel="chrome",
            headless=False,
            args=[
                "--disable-blink-features=AutomationControlled",
                "--no-sandbox",
                "--start-maximized"
            ],
            viewport=None
        )

        page = context.pages[0] if context.pages else context.new_page()

        print(f"[2/6] Membuka URL kuesioner: {target_row['url']}")
        page.goto(target_row["url"], wait_until="domcontentloaded", timeout=60000)
        time.sleep(3.0)

        # Cek jika perlu login SSO BPS
        body_text = page.evaluate("() => document.body ? document.body.innerText : ''")
        if "sso.bps.go.id" in page.url or "/app/login" in page.url or "Selamat Datang Kembali" in body_text or page.locator('input[type="password"]').is_visible():
            print("\n" + "="*60)
            print(">>> [PERHATIAN] Halaman Login SSO terdeteksi!")
            print(">>> Silakan login secara manual di jendela Google Chrome...")
            print("="*60)
            for _ in range(120): # Menunggu hingga 4 menit
                time.sleep(2)
                if "sso.bps.go.id" not in page.url and "/app/login" not in page.url:
                    print(">>> [LOGIN SUKSES] Terdeteksi telah masuk! Menunggu sesi aktif...")
                    time.sleep(3)
                    break

        # Jika pasca login diarahkan ke /app/surveys, buka kembali URL edit
        if "/app/surveys" in page.url:
            print(">>> Membuka ulang link formulir kuesioner...")
            page.goto(target_row["url"], wait_until="domcontentloaded", timeout=60000)
            time.sleep(3.0)

        # Cek proteksi WAF
        body_text = page.evaluate("() => document.body ? document.body.innerText : ''")
        page_title = page.title()
        if "Bot Detected" in page_title or "BOT-" in body_text or "perilaku yang tidak wajar" in body_text:
            print("\n[PERINGATAN WAF] Server BPS mendeteksi Bot! Proses dihentikan untuk keamanan akun.")
            context.close()
            return

        # [3/6] Navigasi ke BLOK II
        print("[3/6] Mengakses menu SE2026 - L BLOK II...")
        time.sleep(2.0)
        try:
            blok2_btn = page.get_by_text("SE2026 - L BLOK II", exact=True)
            blok2_btn.click(timeout=8000)
        except Exception:
            page.evaluate("""() => {
                const el = Array.from(document.querySelectorAll('div, a, button, span')).find(e => e.innerText && e.innerText.trim().includes('SE2026 - L BLOK II'));
                if (el) el.click();
            }""")
        time.sleep(2.5)

        # Cari kartu unit usaha dan klik 'Lihat Rincian'
        keyword = clean_name(target_row["nama_usaha"])
        print(f"      Mencari unit usaha '{target_row['nama_usaha']}' (kata kunci: {keyword})...")
        card_clicked = page.evaluate("""(keyword) => {
            const cards = Array.from(document.querySelectorAll('div, section, tr, li'));
            for (const c of cards) {
                if (c.innerText && c.innerText.toUpperCase().includes(keyword.toUpperCase())) {
                    const btns = Array.from(c.querySelectorAll('button')).filter(b => b.innerText && b.innerText.includes('Lihat Rincian'));
                    if (btns.length > 0) {
                        btns[0].click();
                        return true;
                    }
                }
            }
            const allBtns = Array.from(document.querySelectorAll('button')).filter(b => b.innerText && b.innerText.includes('Lihat Rincian'));
            if (allBtns.length > 0) {
                allBtns[allBtns.length - 1].click();
                return true;
            }
            return false;
        }""", keyword)
        time.sleep(2.5)

        # [4/6] Membuka Catatan pada 13.g
        print("[4/6] Membuka balon Catatan pada Rincian 13. g. Kode KBLI...")
        page.evaluate("""() => {
            const headers = Array.from(document.querySelectorAll('*')).filter(e => e.innerText && e.innerText.trim().startsWith('13. g. Kode KBLI'));
            const targetHeader = headers[headers.length - 1];
            if (!targetHeader) return false;
            let curr = targetHeader;
            while (curr && curr !== document.body) {
                const btns = Array.from(curr.querySelectorAll('button[title="Catatan"]')).filter(b => b.offsetParent !== null || window.getComputedStyle(b).display !== 'none');
                if (btns.length > 0) {
                    btns[0].click();
                    return true;
                }
                curr = curr.parentElement;
            }
            return false;
        }""")
        time.sleep(2.0)

        # Cek pop-up Catatan
        dialog = page.locator('[role="dialog"]')
        if dialog.is_visible():
            dialog_text = dialog.inner_text()
            if "#DC_00" not in dialog_text:
                print(f"      Mengisi catatan '{target_row['status']}'...")
                textarea = dialog.locator('textarea[placeholder="Tambah catatan di sini..."]')
                textarea.fill(target_row["status"])
                time.sleep(1.0)
                save_btn = dialog.locator('button[title="Simpan"]')
                save_btn.click()
                print("      Catatan berhasil disimpan ke daftar aktif.")
                time.sleep(2.0)
            else:
                print("      Catatan #DC_00 sudah ada sebelumnya, melewati pengetikan ulang.")

            # Tutup pop-up catatan dengan Escape dan fallback tombol close
            page.keyboard.press("Escape")
            time.sleep(1.5)
            if dialog.is_visible():
                close_btn = dialog.locator('button[aria-label="Dismiss"], button:has-text("Close"), button[aria-label="Close"], button[title="Tutup"]').first
                if close_btn.is_visible():
                    close_btn.click()
                time.sleep(1.5)

        # [5/6] Klik Kirim di Header & Penanganan Galat
        print("[5/6] Menekan tombol 'Kirim' di sudut kanan atas...")
        kirim_clicked = page.evaluate("""() => {
            const btns = Array.from(document.querySelectorAll('button')).filter(b => (b.innerText || '').trim().includes('Kirim'));
            for (const b of btns) {
                if (!b.closest('[role="dialog"]') && (b.offsetParent !== null || window.getComputedStyle(b).display !== 'none')) {
                    b.click();
                    return true;
                }
            }
            if (btns.length > 0) {
                btns[btns.length - 1].click();
                return true;
            }
            return false;
        }""")
        time.sleep(3.0)

        # Periksa modal Galat
        dialog = page.locator('[role="dialog"]')
        if dialog.is_visible():
            dialog_txt = dialog.inner_text()
            if "GALAT" in dialog_txt or "Perlu diperbaiki" in dialog_txt:
                print("      [GALAT TERDETEKSI] Menangani switch wajib di menu CATATAN...")
                batal_btn = dialog.locator('button:has-text("Batal"), button:has-text("Close")').first
                if batal_btn.is_visible():
                    batal_btn.click()
                time.sleep(1.5)

                # Buka menu CATATAN
                page.locator('aside, nav, div').locator('text="CATATAN"').last.click()
                time.sleep(2.0)

                # Aktifkan toggle Tampilkan Anomali
                page.evaluate("""() => {
                    const ctrl = document.querySelector('div[id$="-control"], [role="switch"]');
                    if (ctrl) ctrl.click();
                }""")
                time.sleep(1.5)

                # Klik Kirim kembali
                page.evaluate("""() => {
                    const btns = Array.from(document.querySelectorAll('button')).filter(b => (b.innerText || '').trim().includes('Kirim'));
                    if (btns.length > 0) btns[btns.length - 1].click();
                }""")
                time.sleep(2.5)

        # Klik tombol Kirim di dalam dialog ringkasan
        dialog_after = page.locator('[role="dialog"]')
        if dialog_after.is_visible():
            modal_kirim_btn = dialog_after.locator('button:has-text("Kirim")').last
            if modal_kirim_btn.is_visible():
                modal_kirim_btn.click()
                time.sleep(2.5)

        # Konfirmasi Pengiriman
        confirm_dialog = page.locator('[role="dialog"]')
        if confirm_dialog.is_visible():
            confirm_btn = confirm_dialog.locator('button:has-text("Konfirmasi")').first
            if confirm_btn.is_visible():
                print("      Menekan tombol 'Konfirmasi' pengiriman...")
                confirm_btn.click()
                time.sleep(3.5)

        # [6/6] Verifikasi & Simpan Excel
        print("[6/6] Memperbarui status di file Excel...")
        page_body = page.evaluate("() => document.body ? document.body.innerText : ''")
        if not ("berhasil" in page_body.lower() or "/app/surveys" in page.url):
            print("[GAGAL] Tidak ada konfirmasi sukses dari FASIH. Status TIDAK ditandai TERKIRIM.")
            context.close()
            return
        ws.cell(row=target_row["row_idx"], column=col_submit, value="TERKIRIM")
        try:
            wb.save(EXCEL_PATH)
            print(f"[SUKSES] Baris {target_row['row_idx']} ({target_row['nama_usaha']}) berhasil di-submit dan ditandai 'TERKIRIM' di {EXCEL_PATH}!")
        except PermissionError:
            backup_path = r"D:\Anomali\Anomaliku_autosave.xlsx"
            print(f"[PERINGATAN] {EXCEL_PATH} terkunci di Excel, menyimpan ke {backup_path}...")
            wb.save(backup_path)
            print(f"[SUKSES] Berhasil disimpan ke {backup_path}!")

        time.sleep(3.0)
        context.close()

    print("\n=== UJI COBA 1 SAMPEL SELESAI DENGAN SUKSES ===")

if __name__ == "__main__":
    run_single_test()
